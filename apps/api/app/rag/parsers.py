import csv
import io
import zipfile
from html.parser import HTMLParser

from app.rag.text import normalize_text
from app.services.errors import ValidationFailed

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".html", ".htm"}


class UnparseableDocument(Exception):
    """File looked valid but could not be parsed; recorded as a failed ingestion job."""


def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    """Cheap pre-checks that reject a file outright (no ingestion record is created)."""
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise ValidationFailed(f"Unsupported file type; allowed: {sorted(ALLOWED_EXTENSIONS)}")
    if not data:
        raise ValidationFailed("File is empty")
    if len(data) > max_bytes:
        raise ValidationFailed(f"File exceeds the {max_bytes // (1024 * 1024)} MB limit")
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValidationFailed("File content is not a PDF")
    if ext == ".docx" and not (data.startswith(b"PK") and zipfile.is_zipfile(io.BytesIO(data))):
        raise ValidationFailed("File content is not a DOCX")
    if ext not in {".pdf", ".docx"} and b"\x00" in data[:4096]:
        raise ValidationFailed("File appears to be binary, not text")
    return ext


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    raise UnparseableDocument("Could not decode text")


class _Html(HTMLParser):
    _SKIP = {"script", "style", "noscript", "template"}
    _BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "h1", "h2", "h3", "h4",
              "h5", "h6", "table", "ul", "ol"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")
            if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                self.parts.append("#" * int(tag[1]) + " ")
            elif tag == "li":
                self.parts.append("- ")

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip:
            self._skip -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def parse_document(ext: str, data: bytes) -> str:
    try:
        if ext in {".txt", ".md"}:
            text = _decode(data)
        elif ext in {".html", ".htm"}:
            p = _Html()
            p.feed(_decode(data))
            text = "".join(p.parts)
        elif ext == ".csv":
            rows = list(csv.reader(io.StringIO(_decode(data))))
            if not rows:
                raise UnparseableDocument("CSV has no rows")
            header, body = rows[0], rows[1:]
            lines = ["; ".join(f"{h.strip()}: {v.strip()}" for h, v in zip(header, r, strict=False)
                               if v.strip()) for r in body]
            text = "\n".join(line for line in lines if line)
        elif ext == ".docx":
            import docx

            d = docx.Document(io.BytesIO(data))
            out = []
            for para in d.paragraphs:
                style = (para.style.name or "") if para.style is not None else ""
                if para.text.strip() and style.startswith("Heading"):
                    level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
                    out.append("#" * level + " " + para.text.strip())
                elif para.text.strip():
                    out.append(para.text.strip())
            for table in d.tables:
                for row in table.rows:
                    out.append("; ".join(c.text.strip() for c in row.cells if c.text.strip()))
            text = "\n".join(out)
        elif ext == ".pdf":
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise UnparseableDocument("Encrypted PDFs are not supported")
            text = "\n\n".join((pg.extract_text() or "") for pg in reader.pages)
        else:  # pragma: no cover - guarded by validate_upload
            raise UnparseableDocument(f"No parser for {ext}")
    except UnparseableDocument:
        raise
    except Exception as exc:  # parser libraries raise many types on corrupt input
        raise UnparseableDocument(f"Could not parse {ext} file: {type(exc).__name__}") from exc
    text = normalize_text(text)
    if not text:
        raise UnparseableDocument("No extractable text (scanned/image-only documents need OCR)")
    return text
